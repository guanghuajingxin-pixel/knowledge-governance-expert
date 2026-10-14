"""FAQ 批量导入解析：CSV / Excel 模板 → 结构化问答行。

模板列（兼容中英文表头，列序不限）：
  问题 | 相似问 | 答案 | 关键词

相似问语义：
- 单元格内多个相似问用分号（;；）或换行分隔
- 同一问题拆成多行（问题与答案均相同）时自动合并为一条问答，相似问聚合去重；
  相似问始终归属主问题，FAQ 列表中仅算一条数据
"""
import io, csv, re
from openpyxl import load_workbook

# 相似问单元格内分隔符：分号（半/全角）与换行；逗号可能出现在问句文本中，不作为分隔符
_SIMILAR_SPLIT = re.compile(r"[;；\n]+")

# 表头别名 → 标准字段
_HEADER_ALIASES = {
    "question": "question", "问题": "question", "q": "question",
    "similar_questions": "similar", "相似问": "similar", "similar": "similar",
    "answer": "answer", "答案": "answer", "a": "answer",
    "keywords": "keywords", "关键词": "keywords", "keyword": "keywords",
}


def _cell(v) -> str:
    return str(v).strip() if v is not None else ""


def _map_headers(headers) -> dict[str, str]:
    """表头行 → {列号: 标准字段}；不认识的列忽略。"""
    mapped = {}
    for i, h in enumerate(headers):
        key = _HEADER_ALIASES.get(_cell(h).lower())
        if key:
            mapped[i] = key
    return mapped


def _parse_similar(raw: str) -> list[str]:
    return [s.strip() for s in _SIMILAR_SPLIT.split(raw) if s.strip()]


def _row(mapped: dict[str, str], values) -> dict | None:
    rec = {}
    for i, key in mapped.items():
        if i < len(values):
            rec[key] = _cell(values[i])
    question = rec.get("question", "")
    answer = rec.get("answer", "")
    if not question or not answer:
        return None
    similar = _parse_similar(rec.get("similar", ""))
    if question in similar:
        similar = [s for s in similar if s != question]
    kw_raw = rec.get("keywords", "")
    keywords = [k.strip() for k in kw_raw.replace("，", ",").split(",") if k.strip()]
    return {"question": question, "answer": answer,
            "similar_questions": similar, "keywords": keywords}


def _decode_csv(data: bytes) -> str:
    """CSV 文本解码：优先 UTF-8（带 BOM），回退 GB18030（Windows Excel 默认另存编码）。"""
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError("CSV 编码无法识别，请用 Excel/WPS 另存为 .xlsx，或以 UTF-8 编码保存 CSV 后重试")


async def parse(file) -> tuple[list[dict], int]:
    """解析上传文件并按（问题, 答案）合并同问题的多行。

    返回 (合并后的问答行列表, 解析无效行数)；无效行 = 缺问题或缺答案。
    所有文件级异常（编码/格式/损坏）统一抛 ValueError，由路由转 400 友好提示。
    """
    data = await file.read()
    name = (file.filename or "").lower()
    rows: list[dict | None] = []
    if name.endswith(".csv"):
        reader = csv.reader(io.StringIO(_decode_csv(data)))
        parsed = list(reader)
        if not parsed:
            return [], 0
        mapped = _map_headers(parsed[0])
        rows = [_row(mapped, r) for r in parsed[1:]]
    elif name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data))
            ws = wb.active
            all_rows = list(ws.iter_rows(values_only=True))
        except Exception as exc:
            raise ValueError(
                f"Excel 文件解析失败，请确认为有效的 .xlsx 文件（{type(exc).__name__}）"
            ) from exc
        if not all_rows:
            return [], 0
        mapped = _map_headers(all_rows[0])
        rows = [_row(mapped, r) for r in all_rows[1:]]
    elif name.endswith(".xls"):
        raise ValueError("暂不支持旧版 .xls，请在 Excel/WPS 中另存为 .xlsx 后重新上传")
    else:
        raise ValueError("仅支持 csv/xlsx")

    skipped = sum(1 for r in rows if r is None)
    # 同（问题, 答案）多行合并：相似问聚合去重，关键词并集
    merged: dict[tuple[str, str], dict] = {}
    ordered: list[dict] = []
    for r in rows:
        if not r:
            continue
        key = (r["question"], r["answer"])
        if key not in merged:
            item = dict(r)
            merged[key] = item
            ordered.append(item)
        else:
            item = merged[key]
            for s in r["similar_questions"]:
                if s not in item["similar_questions"]:
                    item["similar_questions"].append(s)
            for k in r["keywords"]:
                if k not in item["keywords"]:
                    item["keywords"].append(k)
    return ordered, skipped
